from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("outputs/Heyelab-SandFest-Mutual-NDA.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 98, 108)
LIGHT = "F2F4F7"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0,0,0)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ("Heading 1",16,BLUE,16,8), ("Heading 2",13,BLUE,12,6), ("Heading 3",12,DARK,8,4)
]:
    st=styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=color
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

header=sec.header.paragraphs[0]
header.text="MUTUAL NON-DISCLOSURE AGREEMENT  |  HEYELAB + SAND FEST"
header.alignment=WD_ALIGN_PARAGRAPH.CENTER
for run in header.runs:
    run.font.name="Calibri"; run.font.size=Pt(8.5); run.font.bold=True; run.font.color.rgb=GRAY

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run("CONFIDENTIAL  |  Signable template - legal review recommended")
r.font.name="Calibri"; r.font.size=Pt(8); r.font.color.rgb=GRAY

title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(4)
r=title.add_run("MUTUAL NON-DISCLOSURE AGREEMENT")
r.bold=True; r.font.name="Calibri"; r.font.size=Pt(22); r.font.color.rgb=RGBColor(0,0,0)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(14)
r=sub.add_run("Local-Only Data Processing and No-Retention Terms")
r.italic=True; r.font.size=Pt(12); r.font.color.rgb=GRAY

p=doc.add_paragraph()
p.add_run("Effective Date: ").bold=True; p.add_run("____________________")
p=doc.add_paragraph()
p.add_run("Parties. ").bold=True
p.add_run("This Mutual Non-Disclosure Agreement (the “Agreement”) is between ")
p.add_run("Heyelab [full legal name and entity type]").bold=True
p.add_run(", with an address at __________________________________________ (“Heyelab”), and ")
p.add_run("Texas SandFest [full legal name and entity type]").bold=True
p.add_run(", with an address at __________________________________________ (“SandFest”). Heyelab and SandFest are each a “Party” and together the “Parties.”")

def heading(text):
    doc.add_paragraph(text, style="Heading 1")

def clause(num, title, body):
    p=doc.add_paragraph()
    p.paragraph_format.keep_together=False
    r=p.add_run(f"{num}. {title}. "); r.bold=True
    p.add_run(body)

clause(1,"Purpose","The Parties may disclose information to evaluate, develop, support, or perform technology, data, website, document-ingestion, operational, or related services for SandFest (the “Purpose”).")
clause(2,"Confidential Information","“Confidential Information” means nonpublic information disclosed by or for a Party in any form, including business plans, financial information, donor, sponsor, vendor, volunteer, employee, attendee, customer, contact, credential, security, operational, technical, source-code, document, database, analytics, and personal information; the existence and terms of nonpublic projects; and notes, outputs, or materials derived from that information. “SandFest Data” means all Confidential Information, files, records, credentials, content, datasets, and personal information supplied by or collected for SandFest.")
clause(3,"Exclusions","Confidential Information does not include information the receiving Party can document: (a) was lawfully public without breach of this Agreement; (b) was already lawfully known without a confidentiality duty; (c) was independently developed without use of the disclosing Party’s Confidential Information; or (d) was lawfully received from a third party without a confidentiality duty.")
clause(4,"Use and Protection","The receiving Party will use Confidential Information only for the Purpose; limit access to personnel who need it for the Purpose and are bound by duties at least as protective as this Agreement; apply reasonable administrative, physical, and technical safeguards; and not disclose it except as this Agreement permits. Each Party remains responsible for its personnel and contractors. SandFest Data is subject to the stricter requirements in Section 5.")

heading("Local-Only Data Commitments")
clause(5,"SandFest Data: Local Processing; No External Transfer; No Retention","Heyelab will process SandFest Data only on machines physically controlled by Heyelab and configured for the Purpose. Heyelab will not transmit, upload, synchronize, replicate, disclose, or otherwise make SandFest Data available to any third party or external system, including cloud storage, hosted databases, external backup services, analytics or telemetry services, software-as-a-service tools, or externally hosted artificial-intelligence or machine-learning services. Heyelab will not use SandFest Data to train or improve any model or generalized product. Heyelab may create only the minimum temporary local working copies technically necessary for the Purpose. Those copies may not be placed in persistent archives or backups and must be securely deleted promptly when the applicable task is complete, when SandFest requests deletion, or upon termination, whichever occurs first. After deletion, Heyelab will not retain SandFest Data, except that Heyelab may retain this Agreement and a deletion confirmation that does not contain SandFest Data. If a technical requirement would conflict with this Section, Heyelab must obtain SandFest’s prior written approval before receiving or processing the affected data.")
clause(6,"Security Incident","Heyelab will notify SandFest without unreasonable delay after discovering unauthorized access to, acquisition of, disclosure of, loss of, or inability to account for SandFest Data. Heyelab will promptly contain and investigate the incident, reasonably cooperate with SandFest, and provide information reasonably needed for SandFest to meet its legal obligations. Heyelab will not notify third parties or regulators on SandFest’s behalf unless legally required or authorized in writing.")
clause(7,"Return, Deletion, and Confirmation","At any time upon SandFest’s written request, and automatically upon termination, Heyelab will stop processing SandFest Data and promptly return it if requested and securely delete all SandFest Data in Heyelab’s possession or control, including temporary local working copies. Upon request, Heyelab will provide a signed written confirmation of completion. Deletion will use reasonable methods appropriate to the storage medium. This Section does not require deletion of data that Heyelab cannot lawfully delete, but any such exception must be identified to SandFest in writing and remains protected and unused.")

heading("General Confidentiality Terms")
clause(8,"Required Disclosure","If law, subpoena, or court order requires disclosure, the receiving Party may disclose only the legally required portion and, if legally permitted, will promptly notify the disclosing Party and reasonably assist its effort to seek protection. Any disclosed information remains Confidential Information for all other purposes.")
clause(9,"Ownership; No License; No Warranty","Confidential Information remains the disclosing Party’s property. No license or ownership right is granted except the limited right to use it for the Purpose. Confidential Information is provided “as is,” without warranty, except as separately agreed in writing.")
clause(10,"Term and Termination","This Agreement begins on the Effective Date. Either Party may terminate it at any time, for any reason or no reason, by written notice effective upon receipt. Termination immediately ends any right to receive or use Confidential Information, but does not affect rights or duties that accrued before termination.")
clause(11,"Survival","The confidentiality and restricted-use duties survive termination for five (5) years after each disclosure. Duties concerning trade secrets survive for as long as the information remains a trade secret under applicable law. Sections 5 through 9 and 12 through 16 survive termination to the extent needed to carry them out.")
clause(12,"Remedies","Unauthorized use or disclosure may cause harm not readily measured by money. A Party may seek temporary, preliminary, or permanent injunctive or equitable relief, in addition to other available remedies, without waiving any defense or requirement imposed by law.")
clause(13,"Governing Law; Venue","Texas law governs this Agreement, without regard to conflict-of-laws rules. The Parties consent to exclusive jurisdiction and venue in the state and federal courts located in __________________ County, Texas. The blank must be completed before signing.")
clause(14,"Notices","Notices under this Agreement must be in writing and delivered personally, by nationally recognized overnight carrier, or by email with confirmation of transmission to the addresses below (or to an updated address given by notice). A termination notice is effective upon receipt.")
clause(15,"No Obligation; Relationship","Neither Party must disclose information or proceed with any transaction or project. This Agreement does not create a partnership, joint venture, agency, fiduciary, employment, or exclusivity relationship.")
clause(16,"Entire Agreement; Amendment; Severability; Waiver; Assignment","This Agreement is the entire agreement on its subject and supersedes prior discussions on confidentiality for the Purpose. Any amendment or waiver must be in a writing signed by both Parties. Failure to enforce a provision is not a waiver. If a provision is unenforceable, it will be limited to the minimum extent necessary and the remainder will continue. Neither Party may assign this Agreement without the other Party’s written consent, except to a successor in a merger or sale of substantially all relevant assets that agrees in writing to be bound.")
clause(17,"Counterparts and Electronic Signatures","The Parties may sign this Agreement in counterparts, including electronic copies and electronic signatures. Each counterpart is an original, and together they form one agreement.")

heading("Signatures")
p=doc.add_paragraph("The individuals signing below represent that they are authorized to bind the named Party.")
p.paragraph_format.space_after=Pt(12)

table=doc.add_table(rows=1, cols=2)
table.autofit=False
table.columns[0].width=Inches(3.25); table.columns[1].width=Inches(3.25)
for cell in table.rows[0].cells:
    cell.width=Inches(3.25); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=OxmlElement('w:tcMar')
    for side,val in [('top','100'),('start','120'),('bottom','100'),('end','120')]:
        e=OxmlElement(f'w:{side}'); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); tcMar.append(e)
    tcPr.append(tcMar)

def signature_cell(cell, party):
    p=cell.paragraphs[0]; r=p.add_run(party); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=BLUE
    fields=["Full legal name:","By (signature):","Printed name:","Title:","Date:","Notice email:"]
    for label in fields:
        p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(2)
        p.add_run(label).bold=True
        p=cell.add_paragraph("________________________________")
        p.paragraph_format.space_after=Pt(2)

signature_cell(table.cell(0,0),"HEYELAB")
signature_cell(table.cell(0,1),"TEXAS SANDFEST")

# Explicit fixed-width table geometry.
tblPr=table._tbl.tblPr
tblW=tblPr.find(qn('w:tblW'))
if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
tblInd=OxmlElement('w:tblInd'); tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa'); tblPr.append(tblInd)
grid=table._tbl.tblGrid
for child in list(grid): grid.remove(child)
for _ in range(2):
    gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),'4680'); grid.append(gc)
for cell in table.rows[0].cells:
    tcW=cell._tc.get_or_add_tcPr().find(qn('w:tcW')); tcW.set(qn('w:w'),'4680'); tcW.set(qn('w:type'),'dxa')

doc.core_properties.title="Mutual Non-Disclosure Agreement - Heyelab and Texas SandFest"
doc.core_properties.subject="Local-only data processing and no-retention confidentiality agreement"
doc.core_properties.author="Heyelab"
doc.core_properties.keywords="NDA, confidentiality, local processing, deletion, Texas SandFest"
doc.save(OUT)
print(OUT.resolve())
